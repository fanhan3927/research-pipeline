"""研发链 → 分发链的反向衔接（PRD 6.4）：把 valuation 算出的退出估值，
写回这个项目的 Pipeline 内部版对应行，并重新生成一份对外版。

**为什么只做这一件事，不做成通用的"改 Pipeline 任意字段"工具**：
入场价格上限是用户的议价底牌，铁律是"任何情况下都不允许写入 Pipeline 的任何版本"。
与其在一个通用编辑脚本里加运行时检查去拦截，不如干脆不给这个能力——本脚本的命令行
参数里根本没有"字段名"这一说，只有 `--exit-valuation` 这一个值可写。这样"入场价格
上限不能进 Pipeline"就不是一条靠脚本逻辑遵守的规则，而是这个工具做不到的事。
如果真的要手动把别的内部字段写进 Pipeline，只能绕开这个脚本手工改，那本身就是一次
明显、可疑、需要三思的动作，不会被日常调用悄悄带过。

用法：
    python scripts/backfill_exit_valuation.py \\
        --pipeline projects/xxx/03_pipeline/Pipeline_内部版.xlsx \\
        --project-name "盛吉盛（宁波）半导体" \\
        --exit-valuation "80-95亿元（2027年科创板，按可比公司PS均值测算）" \\
        [--project-dir projects/xxx]              # 提供则自动写 audit_log.md
        [--skip-external-regen]                    # 默认会顺带重新生成一份对外版
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from audit_log import append_entry  # noqa: E402
from make_external_pipeline import (  # noqa: E402
    PipelineFormatError,
    _default_output_path,
    _find_source_sheet,
    _header_row,
    convert,
)

EXIT_VALUATION_HEADER = "退出估值"
PROJECT_NAME_HEADER = "项目名称"


class ProjectNotFoundError(Exception):
    pass


def backfill(pipeline_path: Path, project_name: str, exit_valuation: str) -> None:
    wb = openpyxl.load_workbook(pipeline_path)
    ws = _find_source_sheet(wb)
    headers = _header_row(ws)

    for required in (PROJECT_NAME_HEADER, EXIT_VALUATION_HEADER):
        if required not in headers:
            raise PipelineFormatError(f"Pipeline 内部版缺少「{required}」列，无法回填")

    name_col = headers[PROJECT_NAME_HEADER]
    exit_col = headers[EXIT_VALUATION_HEADER]

    matched_rows = [
        row for row in ws.iter_rows(min_row=2)
        if row[name_col - 1].value == project_name
    ]
    if not matched_rows:
        existing_names = sorted(
            {str(row[name_col - 1].value) for row in ws.iter_rows(min_row=2) if row[name_col - 1].value}
        )
        raise ProjectNotFoundError(
            f"Pipeline 里找不到项目名「{project_name}」。已有项目：{existing_names}"
        )
    if len(matched_rows) > 1:
        raise ProjectNotFoundError(
            f"Pipeline 里有 {len(matched_rows)} 行项目名都叫「{project_name}」，无法确定回填哪一行，"
            "请先在 Pipeline 里去重或改成唯一名称"
        )

    row = matched_rows[0]
    row[exit_col - 1].value = exit_valuation
    wb.save(pipeline_path)


def _main() -> None:
    parser = argparse.ArgumentParser(description="把退出估值回填进 Pipeline 内部版")
    parser.add_argument("--pipeline", required=True, type=Path, help="Pipeline_内部版.xlsx 路径")
    parser.add_argument("--project-name", required=True, help="要回填的项目名称，须与 Pipeline 里完全一致")
    parser.add_argument("--exit-valuation", required=True, help="退出估值文本（区间+口径，不是单个数字）")
    parser.add_argument("--project-dir", type=Path, default=None, help="项目目录，提供则自动写 audit_log.md")
    parser.add_argument(
        "--skip-external-regen",
        action="store_true",
        help="不自动重新生成对外版（默认会生成，因为退出估值允许出现在对外版）",
    )
    args = parser.parse_args()

    try:
        backfill(args.pipeline, args.project_name, args.exit_valuation)
    except (PipelineFormatError, ProjectNotFoundError) as e:
        print(f"回填失败：{e}", file=sys.stderr)
        sys.exit(1)

    print(f"已回填「{args.project_name}」的退出估值：{args.exit_valuation}")

    if args.project_dir:
        append_entry(
            args.project_dir,
            actor="orchestrator",
            step="回填退出估值",
            detail=f"项目={args.project_name}，退出估值={args.exit_valuation}",
        )

    if args.skip_external_regen:
        return

    output_path = _default_output_path(args.pipeline.parent)
    try:
        warnings, notices = convert(args.pipeline, output_path)
    except PipelineFormatError as e:
        print(f"提醒：退出估值已回填内部版，但重新生成对外版失败：{e}", file=sys.stderr)
        sys.exit(1)

    print(f"对外版已同步更新：{output_path}")
    print("注意：如果之前的对外版已经发出去了，这是一份新版本，仍要走一遍发送前检查，"
          "不要假设内容和上次一样。")
    if notices:
        print("\n结构性提醒：")
        for n in notices:
            print(f"  - {n}")
    if warnings:
        print("\n关键词扫描命中，发送前请人工复核：")
        for w in warnings:
            print(f"  - {w}")

    if args.project_dir:
        detail = f"输出={output_path.name}（因退出估值回填触发的重新生成）"
        if warnings:
            detail += f"；关键词扫描命中 {len(warnings)} 处"
        append_entry(
            args.project_dir,
            actor="orchestrator",
            step="回填退出估值后重新生成对外版",
            detail=detail,
        )
        append_entry(
            args.project_dir,
            actor="user",
            step="发送前检查清单（退出估值更新后）",
            detail="待用户确认：这是退出估值回填后的新版本，需要重新走一遍发送前检查清单",
            pause_type="confirm",
        )


if __name__ == "__main__":
    _main()
